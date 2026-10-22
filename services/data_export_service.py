"""数据导出 service（v2.6.0）

对称数据导入：把数据导出为 Excel

支持：
- 用户 / 项目 / 报销标准（与导入对齐）
- 报销单 / 审计日志（业务数据）
- 自定义筛选
"""
import io
import logging
from typing import List, Dict, Optional
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import and_

from models import db
from models.models import (
    SysUser, ExpenseProject, ExpenseStandard, ExpenseBill,
    PositionLevel, UserProject,
)

log = logging.getLogger(__name__)


# ============= 公共 =============

def _make_xlsx_bytes(sheet_name: str, headers: List[str], rows: List[list]) -> bytes:
    """生成 xlsx bytes"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 表头
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    # 数据
    for row in rows:
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============= 业务导出 =============

def export_users(role: str = None, status: str = 'normal') -> bytes:
    """导出用户"""
    q = SysUser.query
    if role:
        q = q.filter_by(role=role)
    if status:
        q = q.filter_by(account_status=status)
    users = q.order_by(SysUser.id).all()

    headers = ['ID', '用户名', '姓名', '角色', '人员类型', '电话', '邮箱', '部门', '状态', '创建时间']
    rows = []
    for u in users:
        rows.append([
            u.id, u.username, u.real_name, u.role, u.staff_type or '',
            u.phone or '', u.email or '', u.department or '',
            u.account_status or '',
            u.created_at.isoformat() if u.created_at else '',
        ])
    return _make_xlsx_bytes('用户', headers, rows)


def export_projects(status: str = None) -> bytes:
    """导出项目"""
    q = ExpenseProject.query
    if status:
        q = q.filter_by(status=status)
    projects = q.order_by(ExpenseProject.id).all()

    headers = ['ID', '项目编码', '项目名称', '类型', '开始日期', '结束日期', '主审财务', '状态']
    rows = []
    for p in projects:
        fin = SysUser.query.get(p.finance_auditor_id) if p.finance_auditor_id else None
        rows.append([
            p.id, p.project_code, p.project_name, p.project_type or '',
            p.start_date.isoformat() if p.start_date else '',
            p.end_date.isoformat() if p.end_date else '',
            fin.real_name if fin else '',
            p.status,
        ])
    return _make_xlsx_bytes('项目', headers, rows)


def export_standards() -> bytes:
    """导出报销标准"""
    standards = ExpenseStandard.query.order_by(ExpenseStandard.id).all()
    pos_cache = {p.id: p.level_name for p in PositionLevel.query.all()}

    headers = ['ID', '职务等级', '人员类型', '补助类型', '金额（元）', '生效日期', '失效日期', '项目', '备注']
    rows = []
    for s in standards:
        proj = ExpenseProject.query.get(s.project_id) if s.project_id else None
        rows.append([
            s.id, pos_cache.get(s.position_id, ''), s.staff_type,
            s.subsidy_type, float(s.amount or 0),
            s.valid_start.isoformat() if s.valid_start else '',
            s.valid_end.isoformat() if s.valid_end else '',
            proj.project_name if proj else '全局',
            s.remark or '',
        ])
    return _make_xlsx_bytes('报销标准', headers, rows)


def export_bills(
    start_date: date = None,
    end_date: date = None,
    status: str = None,
    project_id: int = None,
    user_id: int = None,
) -> bytes:
    """导出报销单"""
    q = ExpenseBill.query
    if start_date:
        q = q.filter(ExpenseBill.created_at >= start_date)
    if end_date:
        q = q.filter(ExpenseBill.created_at <= end_date)
    if status:
        q = q.filter_by(status=status)
    if project_id:
        q = q.filter_by(project_id=project_id)
    if user_id:
        q = q.filter_by(user_id=user_id)
    bills = q.order_by(ExpenseBill.created_at.desc()).limit(5000).all()

    headers = ['单号', '类型', '报销人', '部门', '项目', '出差起止', '目的地',
               '交通', '住宿', '补助', '其他', '合计', '大写', '状态', '创建时间']
    rows = []
    user_cache = {u.id: u for u in SysUser.query.all()}
    proj_cache = {p.id: p for p in ExpenseProject.query.all()}

    for b in bills:
        user = user_cache.get(b.user_id)
        proj = proj_cache.get(b.project_id)
        rows.append([
            b.bill_no,
            '项目总账' if b.bill_type == 'project_total' else '个人报销',
            user.real_name if user else '',
            user.department if user else '',
            proj.project_name if proj else '',
            f"{b.travel_start_date or ''} → {b.travel_end_date or ''}",
            b.destination or '',
            float(b.transport_fee or 0),
            float(b.accommodation_fee or 0),
            float(b.subsidy_fee or 0),
            float(b.other_fee or 0),
            float(b.total_amount or 0),
            b.capital_amount or '',
            b.status,
            b.created_at.isoformat() if b.created_at else '',
        ])
    return _make_xlsx_bytes('报销单', headers, rows)
