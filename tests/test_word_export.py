"""
Word 制式单据导出测试
"""
import io
import zipfile
from datetime import date, datetime

import pytest
import requests

from models import db
from models.models import SysUser, ExpenseProject, ExpenseBill, ExpenseBillItem
from werkzeug.security import generate_password_hash


def login(base_url, username='admin', password='NewP@ss123'):
    s = requests.Session()
    r = s.post(f'{base_url}/api/auth/login',
               json={'username': username, 'password': password})
    assert r.json()['code'] == 0, r.json()
    return s


@pytest.fixture
def word_bill_data(app, admin_user, project):
    """测试用单据 + 明细（走全量审批字段便于验证审批区）"""
    with app.app_context():
        u = SysUser(
            username='r_word', real_name='王五', role='reimburser',
            staff_type='在编', account_status='normal',
            department='保障部',
            password_hash=generate_password_hash('Test@001'),
        )
        db.session.add(u)
        db.session.flush()

        auditor = SysUser(
            username='a_word', real_name='稽核员乙', role='auditor',
            staff_type='在编', account_status='normal',
            password_hash=generate_password_hash('Test@001'),
        )
        finance = SysUser(
            username='f_word', real_name='财务丙', role='finance',
            staff_type='在编', account_status='normal',
            password_hash=generate_password_hash('Test@001'),
        )
        db.session.add_all([auditor, finance])
        db.session.flush()

        bill = ExpenseBill(
            bill_no='B-WORD-00001', user_id=u.id, project_id=project['id'],
            bill_type='personal',
            travel_start_date=date(2026, 3, 1), travel_end_date=date(2026, 3, 3),
            destination='上海',
            transport_fee=1000, accommodation_fee=800,
            subsidy_fee=300, other_fee=200, total_amount=2300,
            capital_amount='贰仟叁佰元整',
            status='approved', creator_id=u.id,
            auditor_id=auditor.id, auditor_audit_at=datetime(2026, 3, 4, 10, 0),
            auditor_comment='票据齐全，同意报销',
            finance_auditor_id=finance.id, finance_audit_at=datetime(2026, 3, 5, 14, 30),
            finance_comment='终审通过',
        )
        db.session.add(bill)
        db.session.flush()

        items = [
            ExpenseBillItem(bill_id=bill.id, item_type='交通', amount=600,
                            expense_date=date(2026, 3, 1), description='高铁往返'),
            ExpenseBillItem(bill_id=bill.id, item_type='住宿', amount=800,
                            expense_date=date(2026, 3, 2), description='酒店两晚'),
        ]
        db.session.add_all(items)
        db.session.commit()

        return {'bill_id': bill.id, 'user_id': u.id,
                'auditor_id': auditor.id, 'finance_id': finance.id}


def _doc_tables_text(path):
    """用 python-docx 重新打开，收集全部表格文本（按行拼接）"""
    from docx import Document
    doc = Document(path)
    rows = []
    for table in doc.tables:
        for row in table.rows:
            rows.append(' | '.join(c.text.strip() for c in row.cells))
    return rows


# ============= 单元测试（直接调 generate_bill_docx）=============

def test_generate_bill_docx_file(app, admin_user, project, word_bill_data, tmp_path):
    """生成合法 .docx（ZIP 头 + word/document.xml）"""
    from utils.word_export import generate_bill_docx

    with app.app_context():
        bill = ExpenseBill.query.get(word_bill_data['bill_id'])
        items = ExpenseBillItem.query.filter_by(bill_id=bill.id).all()
        user = SysUser.query.get(bill.user_id)
        proj = ExpenseProject.query.get(bill.project_id)
        auditor = SysUser.query.get(bill.auditor_id)
        finance = SysUser.query.get(bill.finance_auditor_id)

        out = str(tmp_path / 'bill.docx')
        generate_bill_docx(bill, items, proj, user, auditor, finance, out)

        with open(out, 'rb') as f:
            data = f.read()
        assert data[:2] == b'PK'  # ZIP 文件头
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            assert 'word/document.xml' in zf.namelist()
            xml = zf.read('word/document.xml').decode('utf-8')
            # 关键制式要素进入文档
            assert '差旅费报销单' in xml
            assert 'B-WORD-00001' in xml
            assert '贰仟叁佰元整' in xml
            assert '上海' in xml


def test_generate_bill_docx_structure(app, admin_user, project, word_bill_data, tmp_path):
    """结构校验：标题 + 表格体系（编号/主表/费用汇总/大写/明细/审批/签字栏）"""
    from utils.word_export import generate_bill_docx
    from docx import Document

    with app.app_context():
        bill = ExpenseBill.query.get(word_bill_data['bill_id'])
        items = ExpenseBillItem.query.filter_by(bill_id=bill.id).all()
        user = SysUser.query.get(bill.user_id)
        proj = ExpenseProject.query.get(bill.project_id)
        auditor = SysUser.query.get(bill.auditor_id)
        finance = SysUser.query.get(bill.finance_auditor_id)

        out = str(tmp_path / 'bill.docx')
        generate_bill_docx(bill, items, proj, user, auditor, finance, out)

        doc = Document(out)
        # 标题段
        titles = [p.text for p in doc.paragraphs if p.text.strip()]
        assert titles[0] == '差旅费报销单'

        rows = _doc_tables_text(out)
        all_text = '\n'.join(rows)
        # 主表信息
        assert any('报销人' in r and '王五' in r for r in rows)
        assert any('出差目的地' in r and '上海' in r for r in rows)
        # 费用汇总
        assert any('费用项目' in r and '交通费' in r for r in rows)
        assert any('2300.00' in r for r in rows)
        # 费用明细（items）
        assert any('高铁往返' in r for r in rows)
        # 审批意见
        assert any('稽核员初审' in r and '票据齐全' in r for r in rows)
        assert any('财务终审' in r and '终审通过' in r for r in rows)
        # 签字栏
        assert any('领款人签字' in r for r in rows)


def test_generate_bill_docx_capital_fallback(app, admin_user, project, tmp_path):
    """capital_amount 为空时回退 amount_to_capital"""
    from utils.word_export import generate_bill_docx
    from utils.bill_number import amount_to_capital

    with app.app_context():
        u = SysUser.query.filter_by(username='admin').first()
        bill = ExpenseBill(
            bill_no='B-WORD-00002', user_id=u.id, project_id=project['id'],
            bill_type='personal',
            travel_start_date=date(2026, 3, 1), travel_end_date=date(2026, 3, 1),
            destination='北京',
            transport_fee=100, total_amount=100,
            capital_amount=None,  # 空 → 回退
            status='draft', creator_id=u.id,
        )
        db.session.add(bill)
        db.session.commit()

        out = str(tmp_path / 'bill2.docx')
        generate_bill_docx(bill, [], ExpenseProject.query.get(project['id']),
                           u, None, None, out)

        expected = amount_to_capital(100.0)
        rows = _doc_tables_text(out)
        assert any(expected in r for r in rows)


def test_generate_bill_docx_agent_filled(app, admin_user, project, tmp_path):
    """代填单据标注"""
    from utils.word_export import generate_bill_docx
    from docx import Document

    with app.app_context():
        u = SysUser.query.filter_by(username='admin').first()
        bill = ExpenseBill(
            bill_no='B-WORD-00003', user_id=None, project_id=project['id'],
            bill_type='project_total',
            travel_start_date=date(2026, 3, 1), travel_end_date=date(2026, 3, 2),
            destination='广州', total_amount=500, capital_amount='伍佰元整',
            status='draft', creator_id=u.id, is_agent_filled=True,
        )
        db.session.add(bill)
        db.session.commit()

        out = str(tmp_path / 'bill3.docx')
        generate_bill_docx(bill, [], ExpenseProject.query.get(project['id']),
                           None, None, None, out)

        doc = Document(out)
        paras = [p.text for p in doc.paragraphs]
        assert any('稽核员代填' in t for t in paras)
        # 项目总账：报销人单元格显示占位
        rows = _doc_tables_text(out)
        assert any('项目总账' in r for r in rows)


# ============= API 端点测试 =============

def test_docx_api_endpoint(live_url, reimburser, project):
    """API 端点：导出 Word 制式单据"""
    s = login(live_url, 'reimb001', 'Reimb@001!')
    r = s.post(f'{live_url}/api/bills/create', json={
        'user_id': reimburser['id'],
        'project_id': project['id'],
        'bill_type': 'personal',
        'travel_start_date': '2026-08-20',
        'travel_end_date': '2026-08-22',
        'destination': '北京',
        'transport_fee': 1000, 'accommodation_fee': 500,
    })
    bill_id = r.json()['data']['id']

    r = s.get(f'{live_url}/api/bills/{bill_id}/docx')
    assert r.status_code == 200
    assert 'wordprocessingml' in r.headers['content-type']
    assert r.content[:2] == b'PK'  # docx = zip
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        assert 'word/document.xml' in zf.namelist()


def test_docx_api_not_found(live_url):
    """不存在的单据 → 404"""
    s = login(live_url)
    r = s.get(f'{live_url}/api/bills/99999/docx')
    assert r.status_code == 404
    assert r.json()['code'] == 1


def test_docx_api_requires_login(live_url):
    """未登录 → 401"""
    r = requests.get(f'{live_url}/api/bills/1/docx')
    assert r.status_code == 401
